"""
Workflow trigger, status, and data-inspection endpoints.

These routes are a thin HTTP layer over the exact same objects the
terminal demos (examples/demo_run.py) use directly: one shared
WorkflowEngine instance (so RunStateStore, and therefore run history,
persists across requests within a single server process), plus
read-only endpoints that peek at the same local files the connectors
write to (data/agentic.db, data/crm.json, data/outbox.jsonl,
data/workbook.xlsx) so a frontend can prove the run produced real,
persisted side effects rather than just console output.
"""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from typing import Any, Dict, List

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

from src.orchestrator.engine import WorkflowEngine
from src.orchestrator.graph import TaskGraph
from src.orchestrator.scheduler import AutomationScheduler
from src.tools.connectors.crm_connector import DEFAULT_STORE_PATH as CRM_PATH
from src.tools.connectors.db_connector import DEFAULT_DB_PATH
from src.tools.connectors.mail_connector import DEFAULT_OUTBOX_PATH
from src.tools.connectors.sheets_connector import DEFAULT_WORKBOOK_PATH

router = APIRouter()

# One engine per server process. This mirrors what demo_run.py does
# (WorkflowEngine() then .run()); the only difference is the instance is
# shared across requests here so GET /runs/{id} can retrieve a run that a
# previous POST /runs kicked off, exactly like RunStateStore is designed for.
engine = WorkflowEngine()

# Fires engine.run(goal) itself, on a schedule, instead of waiting for a
# POST /runs - see orchestrator/scheduler.py. Started/stopped from
# main.py's lifespan so the background thread has the same lifetime as
# the server process.
scheduler = AutomationScheduler(run_fn=engine.run)


class RunRequest(BaseModel):
    goal: str


class AutomationRequest(BaseModel):
    goal: str
    interval_seconds: int


class AutomationEnabledRequest(BaseModel):
    enabled: bool


def _graph_payload(graph: TaskGraph) -> Dict[str, Any]:
    """Task graph plus results/errors — richer than TaskGraph.to_dict()
    (which the mermaid/demo path uses) since the UI needs result payloads
    to show *what* each tool actually did, not just pass/fail status."""
    tasks = []
    for task in graph.topological_order():
        tasks.append(
            {
                **task.to_dict(),
                "result": task.result,
                "error": task.error,
                "attempts": task.attempts,
            }
        )
    return {"goal": graph.goal, "tasks": tasks, "mermaid": graph.render_mermaid()}


def _state_payload(state) -> Dict[str, Any]:
    return {
        "run_id": state.run_id,
        "goal": state.goal,
        "created_at": state.created_at.isoformat(),
        "replan_attempts": state.replan_attempts,
        "log": state.log,
        "verification": {
            "passed": state.verification.passed,
            "issues": state.verification.issues,
        },
        "graph": _graph_payload(state.graph),
    }


@router.post("/runs")
def create_run(req: RunRequest) -> Dict[str, Any]:
    """Runs the full Planner -> Executor -> Critic loop synchronously and
    returns the finished state. Synchronous on purpose: the underlying
    engine.run() call is itself synchronous (see engine.py), and a demo
    UI benefits more from a single simple call than from a fake
    "in-progress" state that doesn't reflect how the engine works."""
    if not req.goal or not req.goal.strip():
        raise HTTPException(status_code=400, detail="goal must not be empty.")
    try:
        state = engine.run(req.goal.strip())
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=500, detail=f"Run failed: {e}") from e
    return _state_payload(state)


@router.get("/runs")
def list_runs() -> Dict[str, Any]:
    run_ids = engine.state_store.list_run_ids()
    runs = [_state_payload(engine.state_store.load(rid)) for rid in run_ids]
    runs.sort(key=lambda r: r["created_at"], reverse=True)
    return {"runs": runs}


@router.get("/runs/{run_id}")
def get_run(run_id: str) -> Dict[str, Any]:
    state = engine.state_store.load(run_id)
    if state is None:
        raise HTTPException(status_code=404, detail=f"No run with id '{run_id}'.")
    return _state_payload(state)


@router.get("/tools")
def list_tools() -> Dict[str, Any]:
    return {"tools": list(engine.executor.tools.keys())}


# ---------------------------------------------------------------------
# Automations: goals that run themselves on a schedule instead of
# waiting for a POST /runs. Each firing lands in the same run history as
# a manual run (scheduler.py calls the same engine.run()), so /runs,
# the task graph, and the log panel all work unchanged for these.
# ---------------------------------------------------------------------


def _automation_payload(automation) -> Dict[str, Any]:
    return {
        **automation.to_dict(),
    }


@router.post("/automations")
def create_automation(req: AutomationRequest) -> Dict[str, Any]:
    try:
        automation = scheduler.create(req.goal, req.interval_seconds)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e)) from e
    return _automation_payload(automation)


@router.get("/automations")
def list_automations() -> Dict[str, Any]:
    return {"automations": [_automation_payload(a) for a in scheduler.list()]}


@router.patch("/automations/{automation_id}")
def set_automation_enabled(automation_id: str, req: AutomationEnabledRequest) -> Dict[str, Any]:
    try:
        automation = scheduler.set_enabled(automation_id, req.enabled)
    except KeyError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    return _automation_payload(automation)


@router.delete("/automations/{automation_id}")
def delete_automation(automation_id: str) -> Dict[str, Any]:
    scheduler.delete(automation_id)
    return {"deleted": automation_id}


@router.post("/automations/{automation_id}/run_now")
def run_automation_now(automation_id: str) -> Dict[str, Any]:
    """Fires the automation immediately (in addition to its normal
    schedule) and returns its updated status once that run finishes -
    engine.run() is synchronous, so this blocks for the run's duration
    exactly like POST /runs does."""
    try:
        automation = scheduler.run_now(automation_id)
    except KeyError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    return _automation_payload(automation)


# ---------------------------------------------------------------------
# Read-only windows into the same local files the connectors write to.
# Purely for the UI to show "this is real state, not a mock" - none of
# these mutate anything.
# ---------------------------------------------------------------------


@router.get("/data/invoices")
def get_invoices() -> Dict[str, Any]:
    if not Path(DEFAULT_DB_PATH).exists():
        return {"rows": []}
    conn = sqlite3.connect(DEFAULT_DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM invoices ORDER BY id")]
    finally:
        conn.close()
    return {"rows": rows}


@router.get("/data/crm")
def get_crm() -> Dict[str, Any]:
    if not Path(CRM_PATH).exists():
        return {"contact": [], "deal": []}
    with open(CRM_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


@router.get("/data/outbox")
def get_outbox() -> Dict[str, Any]:
    path = Path(DEFAULT_OUTBOX_PATH)
    if not path.exists():
        return {"messages": []}
    messages: List[Dict[str, Any]] = []
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                messages.append(json.loads(line))
    messages.reverse()  # most recent first
    return {"messages": messages}


@router.get("/data/workbook")
def get_workbook() -> Dict[str, Any]:
    path = Path(DEFAULT_WORKBOOK_PATH)
    if not path.exists():
        return {"sheets": {}}
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    sheets: Dict[str, Any] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        sheets[name] = rows
    return {"sheets": sheets}
